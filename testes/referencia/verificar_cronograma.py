#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script de validação do cronograma do projeto Fasoft.
Verifica as aulas agendadas para a data atual (ou informada),
e checa a existência de slides, atividades e materiais de apoio.
"""

import os
import sys
import glob
import re
import argparse
import json
from datetime import datetime, date

# Assegurar saída UTF-8 no terminal Windows
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

def find_fasoft_root(start_dir):
    current = os.path.abspath(start_dir)
    while True:
        agents_file = os.path.join(current, "AGENTS.md")
        if os.path.exists(agents_file):
            try:
                with open(agents_file, "r", encoding="utf-8") as f:
                    content = f.read()
                    if "Fasoft" in content:
                        return current
            except Exception:
                pass
        if os.path.basename(current).lower() == "fasoft" and (
            os.path.exists(os.path.join(current, ".agents")) or os.path.exists(os.path.join(current, "2026"))
        ):
            return current
        parent = os.path.dirname(current)
        if parent == current:
            break
        current = parent
    return None

def normalize_text(text):
    if not text:
        return ""
    # Remove acentos para comparação flexível
    mapping = {
        'á': 'a', 'à': 'a', 'ã': 'a', 'â': 'a',
        'é': 'e', 'ê': 'e',
        'í': 'i',
        'ó': 'o', 'ô': 'o', 'õ': 'o',
        'ú': 'u', 'ü': 'u',
        'ç': 'c'
    }
    t = text.lower()
    for k, v in mapping.items():
        t = t.replace(k, v)
    return t

def extract_keywords(text):
    normalized = normalize_text(text)
    # Extrair palavras com 4 ou mais letras ignorando termos comuns
    stopwords = {"para", "como", "sobre", "pela", "pelo", "com", "sem", "das", "dos", "uma", "este", "esta", "aula", "parte", "secao", "secoes"}
    words = re.findall(r"\b[a-z]{4,}\b", normalized)
    return [w for w in words if w not in stopwords]

def parse_date(val, default_year):
    if isinstance(val, (datetime, date)):
        return val.date() if isinstance(val, datetime) else val
    if isinstance(val, str):
        v = val.strip()
        for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v, fmt).date()
            except ValueError:
                pass
        # Formato dia/mes sem ano
        m = re.match(r"^(\d{1,2})/(\d{1,2})$", v)
        if m:
            day, month = int(m.group(1)), int(m.group(2))
            return date(default_year, month, day)
    return None

def check_materials(disc_dir, aula_num, topico):
    materials = {
        "slides": [],
        "slides_divergentes": [],
        "roteiros_resumos": [],
        "atividades": [],
        "atividades_relacionadas": [],
        "codigos": []
    }

    keywords = extract_keywords(topico)

    # 1. Checar pasta '02 - Aulas'
    aulas_dir = os.path.join(disc_dir, "02 - Aulas")
    if os.path.exists(aulas_dir):
        for fname in os.listdir(aulas_dir):
            if fname.startswith("~$"):
                continue
            is_slide = fname.lower().endswith((".pptx", ".ppt", ".pdf", ".key", ".odp"))
            is_doc = fname.lower().endswith((".md", ".docx", ".doc", ".txt"))

            if aula_num and f"Aula {aula_num}" in fname:
                if is_slide:
                    materials["slides"].append(fname)
                elif is_doc:
                    materials["roteiros_resumos"].append(fname)
            else:
                # Checar se existe arquivo com o tópico mesmo com número diferente
                norm_fname = normalize_text(fname)
                matches = sum(1 for kw in keywords if kw in norm_fname)
                if matches >= 2 and is_slide:
                    materials["slides_divergentes"].append(fname)

    # 2. Checar pasta '03 - Atividades e Listas'
    ativ_dir = os.path.join(disc_dir, "03 - Atividades e Listas")
    if os.path.exists(ativ_dir):
        for fname in os.listdir(ativ_dir):
            if fname.startswith("~$"):
                continue
            if aula_num and f"Aula {aula_num}" in fname:
                materials["atividades"].append(fname)
            else:
                norm_fname = normalize_text(fname)
                matches = sum(1 for kw in keywords if kw in norm_fname)
                if matches >= 2:
                    materials["atividades_relacionadas"].append(fname)

    # 3. Checar pasta '06 - Código e Exemplos'
    code_dir = os.path.join(disc_dir, "06 - Código e Exemplos")
    if os.path.exists(code_dir):
        for fname in os.listdir(code_dir):
            if aula_num and f"aula-{aula_num}" in fname.lower():
                materials["codigos"].append(fname)

    return materials

def run_check(base_dir=None, target_date_str=None, as_json=False):
    if not base_dir:
        base_dir = os.getcwd()

    fasoft_root = find_fasoft_root(base_dir)
    if not fasoft_root:
        err_msg = "[ERRO] Este comando só pode ser executado dentro do projeto Fasoft."
        if as_json:
            print(json.dumps({"error": err_msg, "base_dir": base_dir}))
        else:
            print(err_msg)
            print(f"Diretório atual: {base_dir}")
            print("Nenhum arquivo AGENTS.md com identificação Fasoft foi encontrado na árvore.")
        return 1

    try:
        import openpyxl
    except ImportError:
        err_msg = "[ERRO] A biblioteca openpyxl é necessária para ler os cronogramas em .xlsx."
        if as_json:
            print(json.dumps({"error": err_msg}))
        else:
            print(err_msg)
        return 1

    if target_date_str:
        try:
            target_date = datetime.strptime(target_date_str, "%Y-%m-%d").date()
        except ValueError:
            try:
                target_date = datetime.strptime(target_date_str, "%d/%m/%Y").date()
            except ValueError:
                print(f"[ERRO] Formato de data inválido: {target_date_str}. Use AAAA-MM-DD ou DD/MM/AAAA.")
                return 1
    else:
        target_date = date.today()

    all_xlsx = glob.glob(os.path.join(fasoft_root, "**", "*Cronograma*.xlsx"), recursive=True)
    cronogramas = [
        f for f in all_xlsx 
        if "99 - Arquivo" not in f and not os.path.basename(f).startswith("~$")
    ]

    aulas_hoje = []

    for cpath in cronogramas:
        disc_dir = os.path.dirname(os.path.dirname(cpath))
        disc_name = os.path.basename(disc_dir)

        try:
            wb = openpyxl.load_workbook(cpath, data_only=True)
            sheet = wb.active
        except Exception:
            continue

        for r in range(1, min(sheet.max_row + 1, 150)):
            matched = False
            for col_idx in range(1, min(sheet.max_column + 1, 7)):
                cell_val = sheet.cell(r, col_idx).value
                dt = parse_date(cell_val, target_date.year)
                if dt == target_date:
                    matched = True
                    break

            if matched:
                aula_val = sheet.cell(r, 4).value or sheet.cell(r, 3).value or ""
                topico_val = sheet.cell(r, 5).value or ""
                semana_val = sheet.cell(r, 2).value or ""

                aula_str = str(aula_val).strip()
                topico_str = str(topico_val).strip()
                semana_str = str(semana_val).strip()

                # Extrair número da aula se houver
                match_num = re.search(r"Aula\s*(\d+)", aula_str, re.IGNORECASE)
                aula_num = match_num.group(1).zfill(2) if match_num else None

                materials = check_materials(disc_dir, aula_num, topico_str)

                aulas_hoje.append({
                    "disciplina": disc_name,
                    "disciplina_dir": disc_dir,
                    "cronograma_path": cpath,
                    "semana": semana_str,
                    "aula": aula_str,
                    "aula_num": aula_num,
                    "topico": topico_str,
                    "linha": r,
                    "materiais": materials
                })

    if as_json:
        result = {
            "projeto": fasoft_root,
            "data": target_date.isoformat(),
            "quantidade_aulas": len(aulas_hoje),
            "aulas": aulas_hoje
        }
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0

    # Saída amigável e estruturada
    dias_semana = {
        0: "Segunda-feira",
        1: "Terça-feira",
        2: "Quarta-feira",
        3: "Quinta-feira",
        4: "Sexta-feira",
        5: "Sábado",
        6: "Domingo"
    }
    dia_nome = dias_semana.get(target_date.weekday(), "")

    print("=" * 60)
    print("📅 CRONOGRAMA FASOFT — AULAS DE HOJE")
    print(f"Data: {target_date.strftime('%d/%m/%Y')} ({dia_nome})")
    print(f"Projeto: {fasoft_root}")
    print("=" * 60)

    if not aulas_hoje:
        print(f"\n✨ Nenhuma aula agendada nos cronogramas para {target_date.strftime('%d/%m/%Y')}.")
        print("=" * 60)
        return 0

    print(f"\nFoi encontrada {len(aulas_hoje)} disciplina(s) com agendamento para hoje:\n")

    for idx, item in enumerate(aulas_hoje, 1):
        disc = item["disciplina"]
        semana = item["semana"]
        aula = item["aula"]
        aula_num = item["aula_num"]
        topico = item["topico"]
        mat = item["materiais"]

        print(f"[{idx}] 📘 {disc}")
        if semana:
            print(f"    🗓️  {semana} | {aula}")
        else:
            print(f"    🗓️  {aula}")
        print(f"    🎯 Tópico: {topico}")

        # Tipo especial: Feriado / Recesso / SETECH / Sem Aula
        is_sem_aula = any(k in aula.upper() or k in topico.upper() for k in ["FERIADO", "RECESSO", "SETECH", "SEM AULA"])
        if is_sem_aula:
            print("    🌴 Feriado / Recesso / Evento (sem aula regular).")
            print("-" * 60)
            continue

        # Tipo especial: Avaliação
        if "AVALIA" in aula.upper() or "AVALIA" in topico.upper() or "PROVA" in topico.upper():
            print("    📝 Avaliação / Prova programada.")
            print("-" * 60)
            continue

        # Validação de Slide
        if mat["slides"]:
            print(f"    📊 Slide(s): ✅ Presente")
            for s in mat["slides"]:
                print(f"       • {s}")
        elif mat["slides_divergentes"]:
            print(f"    📊 Slide(s): ⚠️  Não encontrado como 'Aula {aula_num}', mas existe arquivo com tema correspondente:")
            for s in mat["slides_divergentes"]:
                print(f"       • {s}  👉 (Necessita renomear/adaptar)")
        else:
            print(f"    📊 Slide(s): ❌ PENDENTE (Nenhum slide encontrado em '02 - Aulas' para Aula {aula_num or '?'})")

        # Roteiros / Resumos
        if mat["roteiros_resumos"]:
            print(f"    📄 Roteiro/Resumo: ✅ Presente")
            for r in mat["roteiros_resumos"]:
                print(f"       • {r}")

        # Validação de Atividade
        if mat["atividades"]:
            print(f"    📝 Atividade: ✅ Presente")
            for a in mat["atividades"]:
                print(f"       • {a}")
        elif mat["atividades_relacionadas"]:
            print(f"    📝 Atividade: 💡 Exercício(s) relacionado(s) ao tema encontrado(s):")
            for a in mat["atividades_relacionadas"]:
                print(f"       • {a}")
        else:
            print(f"    📝 Atividade: ⚪ Nenhuma atividade cadastrada especificamente para Aula {aula_num or '?'}")

        # Códigos de Exemplo
        if mat["codigos"]:
            print(f"    💻 Código/Exemplo: ✅ Presente ({', '.join(mat['codigos'])})")

        print("-" * 60)

    return 0

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Verifica o cronograma das disciplinas no projeto Fasoft.")
    parser.add_argument("--dir", default=None, help="Diretório de referência para localizar o projeto Fasoft")
    parser.add_argument("--data", default=None, help="Data para consulta (AAAA-MM-DD ou DD/MM/AAAA). Padrão: hoje")
    parser.add_argument("--json", action="store_true", help="Retorna a saída formatada em JSON")
    args = parser.parse_args()

    exit_code = run_check(base_dir=args.dir, target_date_str=args.data, as_json=args.json)
    sys.exit(exit_code)
