"""Gera slides lendo os logos do modelo real (somente leitura); a saída fica em memória."""
import io
from pathlib import Path

import pytest
from pptx import Presentation

from acad.geracao.slides_pptx import GeradorSlidesPptx
from acad.perfil.perfil import carregar_perfil
from testes.unidade.test_slides_pptx import CONTEUDO

MODELO = Path(r"G:\Meu Drive\Fasoft\2026\Segundo Semestre\09 - Modelos\Modelo de Slides UniRV")

pytestmark = pytest.mark.drive


def test_slides_com_logos_do_modelo_real():
    if not MODELO.is_dir():
        pytest.skip("modelo de slides ausente")
    apresentacao = Presentation(io.BytesIO(
        GeradorSlidesPptx(identidade=carregar_perfil("fasoft").slides).gerar(CONTEUDO, MODELO)))
    assert (apresentacao.slide_width, apresentacao.slide_height) == (9144000, 5143500)
    assert sum(1 for s in apresentacao.slides[0].shapes if s.shape_type == 13) == 2
    assert sum(1 for s in apresentacao.slides[1].shapes if s.shape_type == 13) == 1
