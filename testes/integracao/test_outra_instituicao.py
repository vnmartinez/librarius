import io
import json
import zipfile

from pptx import Presentation

from acad.cli import main
from acad.mensagens import contem_termo_tecnico

MARCAS = ("UniRV", "RIO VERDE", "Rio Verde", "Engenharia de Software", "Fasoft")


def _rodar(capsys, *argumentos):
    codigo = main([*argumentos, "--json"])
    saida = json.loads(capsys.readouterr().out)
    assert not contem_termo_tecnico(saida["mensagem"]), saida["mensagem"]
    return codigo, saida


def test_professor_de_outra_instituicao(capsys, semestre_neutro, tmp_path):
    raiz = str(semestre_neutro)
    codigo, saida = _rodar(capsys, "perfil", "criar", "--raiz", raiz, "--id", "faculdade-x",
                           "--instituicao", "Faculdade X", "--cor", "7A1F2B", "--rodape", "FACULDADE X")
    assert codigo == 0
    assert _rodar(capsys, "aplicar", saida["proposta"]["id"])[0] == 0
    assert _rodar(capsys, "iniciar", "2027-1", "--raiz", raiz, "--perfil", "faculdade-x", "--confirmar")[0] == 0

    codigo, saida = _rodar(capsys, "perfil", "detectar", "--raiz", raiz)
    assert codigo == 0 and len(saida["dados"]["disciplinas"]) == 2

    codigo, saida = _rodar(capsys, "disciplina", "criar", "ALG", "--pasta", "Algoritmos", "--nome", "Algoritmos",
                           "--carga", "60", "--unidade", "hora_relogio")
    assert codigo == 0 and saida["dados"]["disciplina"]["semestre"] == "2027-1"
    assert _rodar(capsys, "aplicar", saida["proposta"]["id"])[0] == 0

    codigo, saida = _rodar(capsys, "validar", "--raiz", raiz)
    assert codigo == 0, saida
    assert not [a for a in saida["dados"]["achados"]
                if a["nivel"] == "erro" and a["codigo"].startswith(("pasta-", "nome-"))]

    planejamento = semestre_neutro / "Algoritmos" / "planejamento"
    (planejamento / "aulas.yaml").write_text(
        "versao_esquema: 1\ndisciplina: ALG\naulas:\n- id: Aula 01\n  tema: Grafos\n  status: planejada\n"
        "  objetivos: []\n  prerequisitos: []\n  encontros_necessarios: 1\n", encoding="utf-8")
    conteudo = tmp_path / "conteudo.json"
    conteudo.write_text(json.dumps({
        "tema": "Grafos", "professor": "Profa. Exemplo",
        "roteiro": {"objetivo": "Conhecer grafos", "blocos": [{"titulo": "Grafos", "minutos": 50, "texto": "…"}]},
        "slides": {"slides": [
            {"tipo": "capa", "titulo": "Grafos"},
            {"tipo": "topicos", "titulo": "Ideias", "lead": "O essencial", "itens": ["Vértices", "Arestas"]},
        ]},
    }, ensure_ascii=False), encoding="utf-8")
    codigo, saida = _rodar(capsys, "gerar", "aula", "ALG", "--aula", "01", "--conteudo", str(conteudo))
    assert codigo == 0, saida
    assert not any("logotipos" in aviso for aviso in saida["avisos"])
    assert _rodar(capsys, "aplicar", saida["proposta"]["id"])[0] == 0

    pptx = semestre_neutro / "Algoritmos" / "aulas" / "Aula 01 - Grafos.pptx"
    dados = pptx.read_bytes()
    apresentacao = Presentation(io.BytesIO(dados))
    texto = " ".join(s.text_frame.text for slide in apresentacao.slides for s in slide.shapes if s.has_text_frame)
    assert "FACULDADE X" in texto
    assert not any(marca in texto for marca in MARCAS)
    assert not [s for slide in apresentacao.slides for s in slide.shapes if s.shape_type == 13]
    with zipfile.ZipFile(io.BytesIO(dados)) as pacote:
        xml = b"".join(pacote.read(nome) for nome in pacote.namelist() if nome.endswith(".xml"))
    assert b"015537" not in xml and b"7A1F2B" in xml

    codigo, saida = _rodar(capsys, "cronograma", "consultar", "--raiz",
                           str(semestre_neutro / "Algoritmos" / "aulas"), "--data", "2027-03-01")
    assert codigo == 0
