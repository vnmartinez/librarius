import io, struct, zlib
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu

from acad.geracao.slides_pptx import GeradorSlidesPptx, avisos_editoriais
from acad.perfil.perfil import carregar_perfil

IN = 914400
FASOFT = carregar_perfil("fasoft").slides

CONTEUDO = {
    "tema": "Rotações AVL", "aula_id": "Aula 03", "curso": "Estrutura de Dados 2",
    "professor": "Prof. Fulano", "periodo": "Turma A · 2º semestre de 2026",
    "slides": {"slides": [
        {"tipo": "capa", "titulo": "Rotações AVL"},
        {"tipo": "topicos", "titulo": "Casos", "lead": "Quatro casos de desbalanceamento", "itens": ["LL", "RR", "LR", "RL"],
         "nota": {"rotulo": "Lembre:", "texto": "toda rotação preserva a ordem."}, "notas_professor": "Falar 3 min."},
        {"tipo": "codigo", "titulo": "Rotação à direita", "lead": "Cinco linhas resolvem o caso LL",
         "codigo": "// rotação simples\nfunc rotDir(y) {\n  x = y.esq\n}"},
        {"tipo": "tabela", "titulo": "Complexidade", "lead": "Tudo é logarítmico",
         "linhas": [["Operação", "Custo"], ["Busca", "O(log n)"], ["Inserção", "O(log n)"]]},
        {"tipo": "cards", "titulo": "Objetivos", "lead": "Ao final da aula", "itens": [
            {"tag": "1", "titulo": "Reconhecer", "texto": "Os quatro casos."},
            {"tag": "2", "titulo": "Aplicar", "texto": "A rotação certa."}]},
        {"tipo": "fluxo", "titulo": "Passos", "lead": "Do desbalanceamento à correção",
         "itens": [{"t": "Inserir", "d": "Como na BST"}, {"t": "Subir", "d": "Atualizar alturas"}, {"t": "Rotacionar", "d": "Corrigir"}]},
        {"tipo": "nota", "titulo": "Síntese", "lead": "A ideia central", "rotulo": "Fio condutor:", "texto": "Altura sempre logarítmica."},
        {"tipo": "referencias", "itens": ["CORMEN, T. Algoritmos. 3. ed. Rio de Janeiro: Elsevier, 2012."]},
    ]},
}


def _png(caminho: Path, w=4, h=4):
    def bloco(tipo, dados):
        c = struct.pack(">I", len(dados)) + tipo + dados
        return c + struct.pack(">I", zlib.crc32(tipo + dados) & 0xFFFFFFFF)
    linhas = b"".join(b"\x00" + b"\x01\x55\x37" * w for _ in range(h))
    caminho.write_bytes(b"\x89PNG\r\n\x1a\n" + bloco(b"IHDR", struct.pack(">IIBBBBB", w, h, 8, 2, 0, 0, 0))
                        + bloco(b"IDAT", zlib.compress(linhas)) + bloco(b"IEND", b""))


def _pres(template=None, conteudo=CONTEUDO):
    return Presentation(io.BytesIO(GeradorSlidesPptx(identidade=FASOFT).gerar(conteudo, template)))


def _caixas(slide):
    return [(round(s.left / IN, 2), round(s.top / IN, 2), round(s.width / IN, 2), round(s.height / IN, 2), s) for s in slide.shapes]


def _cor(shape):
    return str(shape.fill.fore_color.rgb)


def test_tamanho_16x9_do_padrao():
    p = _pres()
    assert (p.slide_width, p.slide_height) == (Emu(9144000), Emu(5143500))


def test_um_slide_por_item():
    assert len(_pres().slides) == 8


def test_moldura_institucional_do_slide_de_conteudo():
    caixas = _caixas(_pres().slides[1])
    barra = next(c for c in caixas if c[:4] == (0.47, 0.28, 7.07, 0.63))
    assert _cor(barra[4]) == "015537"
    filete = next(c for c in caixas if c[:4] == (0.18, 5.33, 9.35, 0.11))
    assert _cor(filete[4]) == "015537"
    titulo = next(c for c in caixas if c[:4] == (0.63, 0.28, 6.8, 0.63))
    assert titulo[4].text_frame.text == "Casos"
    run = titulo[4].text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 14 and run.font.bold and run.font.name == "Calibri"
    numero = next(c for c in caixas if c[:4] == (8.95, 5.02, 0.55, 0.26))
    assert numero[4].text_frame.text == "2"


def test_capa_segue_a_geometria_do_padrao():
    caixas = _caixas(_pres().slides[0])
    caixa = next(c for c in caixas if c[:4] == (0.7, 0.5, 8.6, 1.78))
    assert _cor(caixa[4]) == "015537"
    textos = {c[:4]: c[4].text_frame.text for c in caixas if c[4].has_text_frame and c[4].text_frame.text}
    assert textos[(0.8, 0.72, 8.4, 0.62)] == "Estrutura de Dados 2"
    assert textos[(0.9, 1.36, 8.2, 0.78)] == "Aula 03 — Rotações AVL"
    assert textos[(0.7, 2.42, 8.6, 0.3)] == "Prof. Fulano"
    assert textos[(0.7, 4.02, 8.6, 0.26)] == "UNIVERSIDADE DE RIO VERDE (UniRV)"
    assert textos[(0.7, 4.26, 8.6, 0.26)] == "Bacharelado em Engenharia de Software"


def test_fontes_so_calibri_e_courier_e_nunca_abaixo_de_9pt():
    for slide in _pres().slides:
        for forma in slide.shapes:
            if not forma.has_text_frame:
                continue
            for par in forma.text_frame.paragraphs:
                for run in par.runs:
                    assert run.font.name in {"Calibri", "Courier New"}, run.text
                    assert run.font.size.pt >= 9, run.text


def test_codigo_em_courier_new_sobre_fundo_escuro():
    caixas = _caixas(_pres().slides[2])
    assert any(c[4].fill.type == 1 and _cor(c[4]) == "10231B" for c in caixas if hasattr(c[4], "fill"))
    runs = [r for c in caixas if c[4].has_text_frame for p in c[4].text_frame.paragraphs for r in p.runs
            if r.font.name == "Courier New"]
    assert runs and runs[0].text == "// rotação simples"
    assert str(runs[0].font.color.rgb) == "8FB79E"


def test_nada_alem_de_y_5_05_exceto_o_filete_e_o_numero():
    for slide in _pres().slides:
        for esq, topo, larg, alt, forma in _caixas(slide):
            if (esq, topo) in {(0.18, 5.33), (8.95, 5.02)}:
                continue
            assert topo + alt <= 5.05 + 1e-6, (forma.name, topo + alt)


def test_tabela_colunas_somam_a_largura():
    forma = next(s for s in _pres().slides[3].shapes if s.has_table)
    assert sum(c.width for c in forma.table.columns) == forma.width


def test_tabela_cabecalho_verde_e_linhas_alternadas():
    tabela = next(s for s in _pres().slides[3].shapes if s.has_table).table
    assert str(tabela.cell(0, 0).fill.fore_color.rgb) == "015537"
    assert str(tabela.cell(1, 0).fill.fore_color.rgb) == "FFFFFF"
    assert str(tabela.cell(2, 0).fill.fore_color.rgb) == "F2F7F4"


def test_notas_do_professor_vao_nas_notas_do_slide():
    assert _pres().slides[1].notes_slide.notes_text_frame.text == "Falar 3 min."


def test_bullets_usam_marcador_quadrado():
    from lxml import etree
    xml = etree.tostring(_pres().slides[1]._element, encoding="unicode")
    assert "▪" in xml and "•" not in xml


def test_logos_so_com_modelo_na_pasta(tmp_path):
    assert not [s for s in _pres(tmp_path).slides[0].shapes if s.shape_type == 13]
    assets = tmp_path / "assets"
    assets.mkdir()
    for nome in ("logo-header-unirv.png", "logo-unirv.png", "logo-fasoft.png"):
        _png(assets / nome)
    p = _pres(tmp_path)
    figuras = [s for s in p.slides[0].shapes if s.shape_type == 13]
    assert [(round(f.left / IN, 2), round(f.top / IN, 2)) for f in figuras] == [(4.02, 3.16), (4.96, 3.29)]
    cabecalho = [s for s in p.slides[1].shapes if s.shape_type == 13]
    assert [(round(c.left / IN, 2), round(c.top / IN, 2), round(c.width / IN, 2)) for c in cabecalho] == [(7.65, 0.28, 1.91)]


def test_leitor_injetado_e_usado(tmp_path):
    lidos = []
    assets = tmp_path / "assets"
    assets.mkdir()
    for nome in ("logo-header-unirv.png", "logo-unirv.png", "logo-fasoft.png"):
        _png(assets / nome)

    def ler(caminho):
        lidos.append(caminho.name)
        return caminho.read_bytes()

    GeradorSlidesPptx(ler, FASOFT).gerar(CONTEUDO, tmp_path)
    assert {"logo-unirv.png", "logo-fasoft.png", "logo-header-unirv.png"} <= set(lidos)


def test_avisos_editoriais():
    itens = [{"tipo": "topicos", "titulo": "A", "itens": list("1234567")},
             {"tipo": "topicos", "titulo": "B", "lead": "x", "itens": ["a"]},
             {"tipo": "codigo", "titulo": "C", "lead": "y", "codigo": "\n".join(["x"] * 20)}]
    texto = " ".join(avisos_editoriais(itens))
    assert "sem linha-guia" in texto and "mais de 6 itens" in texto
    assert "dois slides seguidos" in texto and "mais de 16 linhas" in texto
    assert avisos_editoriais(CONTEUDO["slides"]["slides"]) == []
