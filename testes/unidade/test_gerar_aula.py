from datetime import datetime
from dataclasses import replace

import pytest
import yaml

from acad.arquivos.privacidade import Workspace
from acad.erros import ErroAcad
from acad.geracao.aula import planejar_aula
from acad.geracao.nomes import arquivos_da_aula, slug
from acad.geracao.slides import GeradorSlidesFalso, obter_gerador
from acad.perfil.perfil import carregar_perfil


DISC = "01 - Disc A - 3º Período"
CONTEUDO = {
    "tema": "Rotações AVL",
    "roteiro": {"objetivo": "Aplicar rotações", "blocos": [
        {"titulo": "Revisão", "minutos": 20, "texto": "…"},
        {"titulo": "Rotações", "minutos": 80, "texto": "…"}]},
    "codigo": {"arquivos": {"main.go": "package main\n"}},
    "lista": {"questoes": [{"enunciado": "Insira 3, 2, 1 numa AVL."}]},
    "slides": {"slides": [{"tipo": "capa", "titulo": "Rotações AVL"}]},
}


@pytest.fixture
def disc(semestre_sintetico):
    pasta = semestre_sintetico / DISC
    (pasta / "01 - Planejamento" / "aulas.yaml").write_text(yaml.safe_dump({
        "versao_esquema": 1, "disciplina": "DA", "aulas": [{
            "id": "Aula 03", "tema": "Rotações AVL", "status": "planejada", "objetivos": [],
            "prerequisitos": [], "encontros_necessarios": 1,
        }],
    }, allow_unicode=True), encoding="utf-8")
    return pasta


def _planejar(ws, semestre, disc, perfil):
    return planejar_aula(ws, semestre, disc, "03", CONTEUDO, GeradorSlidesFalso(),
                         datetime(2026, 9, 1).astimezone(), False, perfil=perfil)


def test_modelo_de_slides_com_outro_nome(disc, semestre_sintetico):
    (semestre_sintetico / "09 - Modelos" / "Slides da Casa").mkdir()
    perfil = replace(carregar_perfil("fasoft"), modelo_slides="Slides da Casa")
    ws = Workspace(semestre_sintetico, perfil.exclusao)
    _, avisos = _planejar(ws, semestre_sintetico, disc, perfil)
    assert not any("logotipos" in aviso for aviso in avisos)


def test_sem_modelo_declarado_nao_avisa_logo(disc, semestre_sintetico):
    perfil = replace(carregar_perfil("fasoft"), modelo_slides=None)
    ws = Workspace(semestre_sintetico, perfil.exclusao)
    _, avisos = _planejar(ws, semestre_sintetico, disc, perfil)
    assert not any("logotipos" in aviso for aviso in avisos)


def test_modelo_declarado_ausente_avisa(disc, semestre_sintetico):
    perfil = carregar_perfil("fasoft")
    ws = Workspace(semestre_sintetico, perfil.exclusao)
    _, avisos = _planejar(ws, semestre_sintetico, disc, perfil)
    assert any("logotipos" in aviso for aviso in avisos)


def test_nomes_canonicos():
    nomes = arquivos_da_aula("03", "Rotações AVL: casos?")
    assert nomes["roteiro"] == "02 - Aulas/Aula 03 - Roteiro de Aula - Rotações AVL casos.md"
    assert nomes["codigo"] == "06 - Código e Exemplos/aula-03-rotacoes-avl-casos"
    assert slug("Árvore B+") == "arvore-b"


def test_planeja_arquivos_e_manifesto(semestre_sintetico, disc):
    operacoes, avisos = planejar_aula(Workspace(semestre_sintetico, []), semestre_sintetico,
                                      disc, "03", CONTEUDO, GeradorSlidesFalso(), datetime(2026, 10, 20, 14, 2))
    destinos = {operacao.destino for operacao in operacoes}
    assert f"{DISC}/02 - Aulas/Aula 03 - Rotações AVL.pptx" in destinos
    assert f"{DISC}/06 - Código e Exemplos/aula-03-rotacoes-avl/main.go" in destinos
    aulas = yaml.safe_load(next(op for op in operacoes if op.destino.endswith("aulas.yaml")).conteudo)
    aula = aulas["aulas"][0]
    assert aula["status"] == "rascunho" and aula["manifesto"]["aprovado"] is False
    assert f"{DISC}/01 - Planejamento/aulas.yaml#Aula 03" in aula["manifesto"]["entradas"]


def test_slides_reais_sao_gerados_e_abrem(semestre_sintetico, disc):
    import io
    from pptx import Presentation

    operacoes, avisos = planejar_aula(Workspace(semestre_sintetico, []), semestre_sintetico,
                                      disc, "03", CONTEUDO, obter_gerador(), datetime(2026, 10, 20))
    pptx = next(op for op in operacoes if op.destino.endswith(".pptx"))
    apresentacao = Presentation(io.BytesIO(pptx.conteudo))
    assert (apresentacao.slide_width, apresentacao.slide_height) == (9144000, 5143500)
    assert len(apresentacao.slides) == 1
    assert any("modelo de slides" in aviso.lower() for aviso in avisos)  # sem pasta de modelo, sem logotipos


def test_avisos_editoriais_chegam_ao_professor(semestre_sintetico, disc):
    conteudo = {**CONTEUDO, "slides": {"slides": [
        {"tipo": "capa"}, {"tipo": "topicos", "titulo": "Casos", "itens": ["a"]}]}}
    _, avisos = planejar_aula(Workspace(semestre_sintetico, []), semestre_sintetico,
                              disc, "03", conteudo, obter_gerador(), datetime(2026, 10, 20))
    assert any("sem linha-guia" in aviso for aviso in avisos)


def test_aula_aprovada_recusa(semestre_sintetico, disc):
    arquivo = disc / "01 - Planejamento" / "aulas.yaml"
    documento = yaml.safe_load(arquivo.read_text(encoding="utf-8"))
    documento["aulas"][0]["status"] = "aprovada"
    arquivo.write_text(yaml.safe_dump(documento, allow_unicode=True), encoding="utf-8")
    with pytest.raises(ErroAcad) as erro:
        planejar_aula(Workspace(semestre_sintetico, []), semestre_sintetico,
                      disc, "03", CONTEUDO, GeradorSlidesFalso(), datetime(2026, 10, 20))
    assert erro.value.codigo == "aula-aprovada"


def test_estimativa_de_tempo(semestre_sintetico, disc):
    longo = {**CONTEUDO, "roteiro": {"objetivo": "x", "blocos": [
        {"titulo": "a", "minutos": 140, "texto": "…"}]}}
    (disc / "01 - Planejamento" / "encontros-TA.yaml").write_text(yaml.safe_dump({
        "versao_esquema": 1, "gerado_por": "acad 0.1.0", "gerado_em": "2026-10-01T10:00:00-03:00",
        "entradas": {}, "encontros": [{"id": "E03", "data": "2026-08-10", "inicio": "19:00", "fim": "20:40",
                                    "status": "previsto", "aulas": ["Aula 03"], "funcao": "conteudo", "origem": "horario"}],
    }), encoding="utf-8")
    _, avisos = planejar_aula(Workspace(semestre_sintetico, []), semestre_sintetico,
                              disc, "03", longo, GeradorSlidesFalso(), datetime(2026, 10, 20))
    assert any("140 min" in aviso and "100 min" in aviso for aviso in avisos)


def test_conteudo_invalido(semestre_sintetico, disc):
    with pytest.raises(ErroAcad) as erro:
        planejar_aula(Workspace(semestre_sintetico, []), semestre_sintetico,
                      disc, "03", {"tema": 1}, GeradorSlidesFalso(), datetime(2026, 10, 20))
    assert erro.value.codigo == "esquema-invalido"


def test_cli_gera_proposta_sem_gravar(capsys, semestre_sintetico, disc, tmp_path):
    import json
    from acad.cli import main

    arquivo = tmp_path / "conteudo.json"
    arquivo.write_text(json.dumps(CONTEUDO, ensure_ascii=False), encoding="utf-8")
    codigo = main(["gerar", "aula", "01", "--aula", "03", "--conteudo", str(arquivo),
                   "--raiz", str(semestre_sintetico), "--json"])
    resultado = json.loads(capsys.readouterr().out)
    assert codigo == 0
    assert resultado["proposta"] and any("aulas.yaml" in item["destino"]
                                         for item in resultado["proposta"]["itens"])
    assert any(item["destino"].endswith(".pptx") for item in resultado["proposta"]["itens"])
    assert not list((disc / "02 - Aulas").glob("*.md"))
