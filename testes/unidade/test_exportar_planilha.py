from pathlib import Path

from openpyxl import load_workbook

from acad.cronograma.exportar_planilha import exportar_planilha
from acad.cronograma.importar import importar_cronograma
from acad.cronograma.planilha import ler_cronograma


XLSX = (Path(__file__).resolve().parents[1] / "amostras/fasoft/2026/Segundo Semestre"
        / "04 - Estrutura de Dados 2 - 5º Período/01 - Planejamento/Cronograma - Estrutura de Dados 2.xlsx")
TURMA = {"id": "TA", "disciplina": "ED2", "fuso": "America/Sao_Paulo",
         "horarios": [{"dia_semana": 2, "inicio": "19:00", "fim": "20:40"},
                      {"dia_semana": 5, "inicio": "19:00", "fim": "20:40"}]}


def test_planilha_ida_e_volta_preserva_encontros_e_aulas(tmp_path):
    anteriores, doc_aulas = importar_cronograma(ler_cronograma(XLSX.read_bytes(), 2026), TURMA, "ED2")
    aulas = {item["id"]: item for item in doc_aulas["aulas"]}
    conteudo = exportar_planilha("Estrutura de Dados 2", anteriores, aulas)
    arquivo = tmp_path / "cronograma.xlsx"
    arquivo.write_bytes(conteudo)
    aba = load_workbook(arquivo).active
    assert aba["B1"].value == "Estrutura de Dados 2"
    assert [aba.cell(3, coluna).value for coluna in range(2, 6)] == ["Semanas", "Datas", "Aulas", "Tópicos"]
    importados, doc_importado = importar_cronograma(ler_cronograma(conteudo, 2026), TURMA, "ED2")
    campos = lambda encontros: [(e.data, e.aulas, e.funcao, e.status) for e in encontros]
    assert campos(importados) == campos(anteriores)
    assert doc_importado == doc_aulas


def test_cli_propõe_substituir_planilha_sem_gravar(capsys, semestre_sintetico):
    import json
    from datetime import datetime
    from acad import esquemas
    from acad.cli import main
    from acad.planejamento.arquivos import documento_encontros

    pasta = semestre_sintetico / "01 - Disc A - 3º Período" / "01 - Planejamento"
    encontros, aulas = importar_cronograma(ler_cronograma(XLSX.read_bytes(), 2026), TURMA, "DA")
    (pasta / "encontros-TA.yaml").write_text(esquemas.para_yaml(
        documento_encontros(encontros, {}, "0.1.0", datetime.now().astimezone())), encoding="utf-8")
    (pasta / "aulas.yaml").write_text(esquemas.para_yaml(aulas), encoding="utf-8")
    original = pasta / "Cronograma - Disc A.xlsx"
    original.write_bytes(XLSX.read_bytes())
    codigo = main(["cronograma", "exportar-planilha", "01", "--turma", "TA",
                   "--raiz", str(semestre_sintetico), "--json"])
    resultado = json.loads(capsys.readouterr().out)
    assert codigo == 0
    assert resultado["proposta"]["itens"][0]["tipo"] == "substituir"
    assert original.read_bytes() == XLSX.read_bytes()
