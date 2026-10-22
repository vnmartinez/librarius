from pathlib import Path

import openpyxl

AMOSTRA = Path(__file__).resolve().parents[1] / "amostras" / "fasoft"
SEMESTRE = AMOSTRA / "2026" / "Segundo Semestre"
ED2 = SEMESTRE / "04 - Estrutura de Dados 2 - 5º Período"


def test_estrutura_espelhada():
    assert (AMOSTRA / "AGENTS.md").read_text(encoding="utf-8").count("Fasoft") >= 1
    assert (ED2 / "02 - Aulas" / "Aula 01 - Apresentação da Disciplina.pptx").exists()
    assert (ED2 / "01 - Planejamento" / "Cronograma - Estrutura de Dados 2.xlsx.bak").exists()


def test_espelho_sem_conteudo_exceto_cronograma():
    for caminho in ED2.rglob("*"):
        if caminho.is_file() and caminho.name != "Cronograma - Estrutura de Dados 2.xlsx":
            assert caminho.stat().st_size == 0, caminho


def test_cronograma_real_legivel():
    pasta = openpyxl.load_workbook(
        ED2 / "01 - Planejamento" / "Cronograma - Estrutura de Dados 2.xlsx", data_only=True
    )
    planilha = pasta.active
    assert planilha.cell(3, 3).value == "Datas"
    assert planilha.cell(4, 4).value == "Aula 01"


def test_nenhum_dado_de_aluno_nem_logo():
    nomes = [caminho.as_posix() for caminho in AMOSTRA.rglob("*")]
    assert not any("Notas" in nome or "02 - Alunos" in nome or "Relatórios" in nome for nome in nomes)
    assert not any(nome.lower().endswith((".png", ".jpg")) for nome in nomes)


def test_oraculo_copiado():
    assert (Path(__file__).resolve().parents[1] / "referencia" / "verificar_cronograma.py").exists()
