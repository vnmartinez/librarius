import io

from pptx import Presentation

from acad.geracao.identidade import IDENTIDADE_NEUTRA, identidade_de, mistura, paleta_de_cor
from acad.geracao.slides_pptx import GeradorSlidesPptx, avisos_editoriais
from acad.perfil.perfil import carregar_perfil
from testes.unidade.test_slides_pptx import CONTEUDO, _png


def _textos(dados):
    p = Presentation(io.BytesIO(dados))
    return " ".join(s.text_frame.text for slide in p.slides for s in slide.shapes if s.has_text_frame)


def test_fasoft_traz_a_identidade_unirv():
    ident = carregar_perfil("fasoft").slides
    assert ident.paleta["principal"] == "015537" and ident.paleta["codigo_fundo"] == "10231B"
    assert ident.capa_linhas == ("UNIVERSIDADE DE RIO VERDE (UniRV)", "Bacharelado em Engenharia de Software")
    assert ident.capa_local == "RIO VERDE — {ano}"
    assert ident.logo_cabecalho == "logo-header-unirv.png"
    assert ident.logos_capa == ("logo-unirv.png", "logo-fasoft.png")


def test_neutra_nao_tem_marca(tmp_path):
    (tmp_path / "assets").mkdir()
    for nome in ("logo-header-unirv.png", "logo-unirv.png", "logo-fasoft.png"):
        _png(tmp_path / "assets" / nome)
    dados = GeradorSlidesPptx().gerar({**CONTEUDO, "ano": 2027}, tmp_path)
    texto = _textos(dados)
    assert "UniRV" not in texto and "RIO VERDE" not in texto and "Engenharia de Software" not in texto
    assert "2027" in texto
    p = Presentation(io.BytesIO(dados))
    assert not [s for slide in p.slides for s in slide.shapes if s.shape_type == 13]
    assert b"015537" not in b"".join(part.blob for part in p.part.package.iter_parts())


def test_bloco_parcial_completa_com_neutra():
    ident = identidade_de({"paleta": {"principal": "AA0000"}, "capa": {"linhas": ["Faculdade X"]}})
    assert ident.paleta["principal"] == "AA0000"
    assert ident.paleta["corpo"] == IDENTIDADE_NEUTRA.paleta["corpo"]
    assert ident.capa_linhas == ("Faculdade X",) and ident.capa_local == "{ano}"


def test_paleta_de_cor():
    assert mistura("000000", "FFFFFF", 0.5) == "808080"
    paleta = paleta_de_cor("1F3A5F")
    assert paleta["principal"] == "1F3A5F"
    assert paleta["claro"] == mistura("1F3A5F", "FFFFFF", 0.88)
    assert paleta["codigo_fundo"] == IDENTIDADE_NEUTRA.paleta["codigo_fundo"]


def test_limites_do_perfil_nos_avisos():
    slides = [{"tipo": "topicos", "titulo": "T", "lead": "l", "itens": ["a", "b", "c"]}]
    assert avisos_editoriais(slides) == []
    assert "mais de 2 itens" in avisos_editoriais(slides, {"topicos": 2, "linhas_codigo": 16, "linhas_tabela": 8})[0]
