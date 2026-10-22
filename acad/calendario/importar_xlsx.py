from __future__ import annotations

import hashlib
import io
from datetime import date, datetime, timedelta

import openpyxl

from acad.arquivos.caminhos import sem_acento
from acad.calendario.modelo import Calendario, DiaEspecial, Evento
from acad.erros import ErroAcad

TIPOS = {
    "feriado": "feriado", "recesso": "recesso", "sabado letivo": "sabado-letivo",
    "troca de dia": "troca-de-dia", "semana academica": "semana-academica",
    "avaliacao": "avaliacao", "outro": "outro",
}
DIAS = {nome: indice for indice, nome in enumerate(
    ("segunda", "terca", "quarta", "quinta", "sexta", "sabado", "domingo"), 1
)}


def _data(valor) -> date:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    for formato in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(valor).strip(), formato).date()
        except ValueError:
            pass
    raise ErroAcad("data-invalida", "Não entendi uma data na planilha do calendário.")


def importar_xlsx(conteudo: bytes, semestre: str, inicio: date, fim: date) -> Calendario:
    planilha = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True).active
    cabecalho = [sem_acento(str(planilha.cell(1, coluna).value or ""))
                 for coluna in range(1, 5)]
    if cabecalho != ["data", "tipo", "descricao", "equivale a"]:
        raise ErroAcad("planilha-sem-cabecalho", "Não reconheci as colunas da planilha do calendário.")
    dias = []
    eventos = []
    for valores in planilha.iter_rows(min_row=2, max_col=4, values_only=True):
        data_bruta, tipo_bruto, descricao_bruta, equivale_bruto = valores
        if data_bruta is None and tipo_bruto is None:
            continue
        partes = str(data_bruta).split(" a ") if isinstance(data_bruta, str) else [data_bruta]
        primeira, ultima = _data(partes[0]), _data(partes[-1])
        tipo = TIPOS.get(sem_acento(str(tipo_bruto or "")))
        if tipo is None or ultima < primeira:
            raise ErroAcad("calendario-invalido", "Há um tipo ou intervalo inválido na planilha do calendário.")
        descricao = str(descricao_bruta or "")
        equivale = DIAS.get(sem_acento(str(equivale_bruto or "")))
        if tipo in ("feriado", "recesso"):
            atual = primeira
            while atual <= ultima:
                dias.append(DiaEspecial(atual, tipo, descricao))
                atual += timedelta(days=1)
        elif primeira < ultima:
            eventos.append(Evento(primeira, ultima, tipo, descricao))
        else:
            dias.append(DiaEspecial(primeira, tipo, descricao, equivale))
    return Calendario(
        semestre, inicio, fim, "rascunho", dias, eventos,
        {"tipo": "xlsx", "sha256": "sha256:" + hashlib.sha256(conteudo).hexdigest()},
    )
