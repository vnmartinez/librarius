from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime

import openpyxl

from acad.arquivos.caminhos import sem_acento
from acad.erros import ErroAcad


@dataclass
class LinhaCronograma:
    linha: int
    semana: str
    data: date | None
    aula: str
    aula_num: str | None
    topico: str


def parse_data(valor, ano_padrao: int) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, str):
        texto = valor.strip()
        for formato in ("%d/%m/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(texto, formato).date()
            except ValueError:
                pass
        correspondencia = re.match(r"^(\d{1,2})/(\d{1,2})$", texto)
        if correspondencia:
            try:
                return date(ano_padrao, int(correspondencia.group(2)), int(correspondencia.group(1)))
            except ValueError:
                return None
    return None


def _num(aula: str) -> str | None:
    correspondencia = re.search(r"Aula\s*(\d+)", aula, re.IGNORECASE)
    return correspondencia.group(1).zfill(2) if correspondencia else None


def _txt(valor) -> str:
    return str(valor).strip() if valor is not None else ""


def _aba(conteudo: bytes):
    return openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True).active


def linhas_na_data(conteudo: bytes, alvo: date) -> list[LinhaCronograma]:
    planilha = _aba(conteudo)
    saida = []
    for linha in range(1, min(planilha.max_row + 1, 150)):
        if any(parse_data(planilha.cell(linha, coluna).value, alvo.year) == alvo
               for coluna in range(1, min(planilha.max_column + 1, 7))):
            aula = _txt(planilha.cell(linha, 4).value or planilha.cell(linha, 3).value or "")
            saida.append(LinhaCronograma(
                linha, _txt(planilha.cell(linha, 2).value), alvo,
                aula, _num(aula), _txt(planilha.cell(linha, 5).value),
            ))
    return saida


def ler_cronograma(conteudo: bytes, ano_padrao: int) -> list[LinhaCronograma]:
    planilha = _aba(conteudo)
    cabecalho = None
    for linha in range(1, min(planilha.max_row, 30) + 1):
        nomes = {
            sem_acento(_txt(planilha.cell(linha, coluna).value)): coluna
            for coluna in range(1, planilha.max_column + 1)
        }
        if "datas" in nomes and "aulas" in nomes:
            cabecalho = (linha, nomes)
            break
    if not cabecalho:
        raise ErroAcad(
            "planilha-sem-cabecalho",
            "Não reconheci as colunas da planilha do cronograma (Datas, Aulas, Tópicos).",
        )
    primeira, nomes = cabecalho
    coluna_data, coluna_aula = nomes["datas"], nomes["aulas"]
    coluna_topico, coluna_semana = nomes.get("topicos"), nomes.get("semanas")
    saida = []
    semana = ""
    for linha in range(primeira + 1, planilha.max_row + 1):
        data = parse_data(planilha.cell(linha, coluna_data).value, ano_padrao)
        aula = _txt(planilha.cell(linha, coluna_aula).value)
        if data is None and not aula:
            break
        if coluna_semana and planilha.cell(linha, coluna_semana).value:
            semana = _txt(planilha.cell(linha, coluna_semana).value)
        saida.append(LinhaCronograma(
            linha, semana, data, aula, _num(aula),
            _txt(planilha.cell(linha, coluna_topico).value) if coluna_topico else "",
        ))
    return saida
