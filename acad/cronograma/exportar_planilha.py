from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from acad.encontros.gerar import Encontro


def exportar_planilha(titulo: str, encontros: list[Encontro],
                      aulas: dict[str, dict]) -> bytes:
    livro = Workbook()
    aba = livro.active
    aba.title = "Cronograma"
    aba["B1"] = titulo
    aba["B1"].font = Font(bold=True, size=16)
    for coluna, rotulo in zip(range(2, 6), ("Semanas", "Datas", "Aulas", "Tópicos")):
        celula = aba.cell(3, coluna, rotulo)
        celula.font = Font(bold=True)
        celula.alignment = Alignment(horizontal="center")
    semanas = {}
    for numero, encontro in enumerate(sorted(encontros, key=lambda e: (e.data, e.inicio)), 4):
        chave = encontro.data.isocalendar()[:2]
        if chave not in semanas:
            semanas[chave] = len(semanas) + 1
            aba.cell(numero, 2, f"Semana {semanas[chave]}")
        celula_data = aba.cell(numero, 3, encontro.data)
        celula_data.number_format = "DD/MM/YYYY"
        if encontro.status == "cancelado":
            aba.cell(numero, 4, "SEM AULA")
            aba.cell(numero, 5, encontro.motivo or "SEM AULA")
        elif encontro.funcao == "avaliacao":
            aba.cell(numero, 4, "AVALIAÇÃO")
            aba.cell(numero, 5, "AVALIAÇÃO")
        else:
            aba.cell(numero, 4, ", ".join(encontro.aulas))
            temas = [aulas[id_aula]["tema"] for id_aula in encontro.aulas if id_aula in aulas]
            aba.cell(numero, 5, "; ".join(temas))
    for coluna, largura in {"B": 16, "C": 16, "D": 22, "E": 66}.items():
        aba.column_dimensions[coluna].width = largura
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
